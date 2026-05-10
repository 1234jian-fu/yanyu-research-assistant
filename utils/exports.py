"""Document export helpers shared by multiple Streamlit pages."""

import re
from io import BytesIO

from docx import Document
from docx.shared import RGBColor


def create_docx_with_redlines(original: str, revised: str, metadata: dict) -> bytes:
    doc = Document()
    doc.add_heading(f"学研·修订模式 - {metadata['function']}", 0)
    info = doc.add_paragraph()
    info.add_run(f"板块: {metadata.get('section', '未指定')}\n")
    info.add_run(f"领域: {metadata.get('domain', '未指定')}\n")
    info.add_run(f"时间: {metadata.get('timestamp', '')}\n")
    doc.add_heading("原文", 1)
    doc.add_paragraph(original)
    doc.add_heading("修改后", 1)
    doc.add_paragraph(revised)
    doc.add_heading("修改说明", 2)

    original_words = set(original.lower().split())
    revised_words = set(revised.lower().split())
    added = revised_words - original_words
    removed = original_words - revised_words
    if added:
        p = doc.add_paragraph()
        run = p.add_run("新增词汇: ")
        run.font.color.rgb = RGBColor(0, 128, 0)
        p.add_run(", ".join(list(added)[:20]))
    if removed:
        p = doc.add_paragraph()
        run = p.add_run("删除词汇: ")
        run.font.color.rgb = RGBColor(255, 0, 0)
        p.add_run(", ".join(list(removed)[:20]))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def create_docx(content: str, metadata: dict) -> bytes:
    doc = Document()
    doc.add_heading(f"学研·工科科研助手 - {metadata.get('function', '导出')}", 0)
    info = doc.add_paragraph()
    for key in ("section", "domain", "timestamp"):
        if metadata.get(key):
            info.add_run(f"{key}: {metadata[key]}\n")
    doc.add_heading("处理结果", 1)
    doc.add_paragraph(content)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def build_export_filename(prefix: str, name_hint: str, suffix: str = ".docx") -> str:
    safe_name = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", (name_hint or "document").strip())
    safe_name = safe_name.strip("._") or "document"
    if not safe_name.lower().endswith(suffix):
        safe_name = f"{safe_name}{suffix}"
    return f"{prefix}_{safe_name}"
